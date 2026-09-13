"""
Server-side resume text extraction.

Files are parsed on the backend and only the extracted text is forwarded to
the AI service. Raw uploads are never relayed to the AI provider.
"""

import logging
import os

logger = logging.getLogger(__name__)

# Allowed upload extensions and the largest file we will accept.
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB

# MIME types we consider consistent with the allowed extensions. Browsers are
# inconsistent here, so this is treated as advisory rather than authoritative.
ALLOWED_MIME_PREFIXES = (
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml",
    "application/octet-stream",
    "text/plain",
)


class ResumeParseError(Exception):
    """Raised when an uploaded resume cannot be validated or read."""

    def __init__(self, message, code="resume_invalid"):
        super().__init__(message)
        self.message = message
        self.code = code


class ResumeParser:
    @staticmethod
    def get_extension(filename):
        return os.path.splitext(str(filename or ""))[1].lstrip(".").lower()

    @staticmethod
    def validate(file_obj, filename=None):
        """
        Validate extension, size and (advisory) MIME type before parsing.

        Raises ResumeParseError with a user-safe message on rejection.
        """
        filename = filename or getattr(file_obj, "name", "")
        ext = ResumeParser.get_extension(filename)

        if not ext:
            raise ResumeParseError(
                "The file has no extension. Upload a PDF or DOCX resume.",
                code="resume_no_extension",
            )

        if ext not in ALLOWED_EXTENSIONS:
            raise ResumeParseError(
                f"Unsupported file type '.{ext}'. Upload a PDF or DOCX resume.",
                code="resume_unsupported_type",
            )

        size = getattr(file_obj, "size", None)
        if size is None:
            try:
                current = file_obj.tell()
                file_obj.seek(0, os.SEEK_END)
                size = file_obj.tell()
                file_obj.seek(current)
            except (OSError, AttributeError):
                size = None

        if size is not None:
            if size == 0:
                raise ResumeParseError(
                    "The uploaded file is empty.", code="resume_empty_file"
                )
            if size > MAX_FILE_BYTES:
                limit_mb = MAX_FILE_BYTES // (1024 * 1024)
                raise ResumeParseError(
                    f"The file is too large ({size / (1024 * 1024):.1f} MB). "
                    f"Maximum size is {limit_mb} MB.",
                    code="resume_too_large",
                )

        content_type = getattr(file_obj, "content_type", None)
        if content_type and not str(content_type).startswith(ALLOWED_MIME_PREFIXES):
            logger.warning(
                "Resume upload '%s' has unexpected MIME type '%s'", filename, content_type
            )

        return ext

    @staticmethod
    def extract_text(file_obj, filename=None, validate=True):
        """
        Extract plain text from an uploaded resume.

        Raises ResumeParseError when the file is invalid or yields no text, so
        callers never silently proceed with an empty resume.
        """
        filename = filename or getattr(file_obj, "name", "")
        ext = ResumeParser.validate(file_obj, filename) if validate else \
            ResumeParser.get_extension(filename)

        text = ""
        try:
            try:
                file_obj.seek(0)
            except (OSError, AttributeError):
                pass

            if ext == "pdf":
                import PyPDF2

                reader = PyPDF2.PdfReader(file_obj)
                if getattr(reader, "is_encrypted", False):
                    try:
                        reader.decrypt("")
                    except Exception as exc:
                        raise ResumeParseError(
                            "This PDF is password protected. Upload an unprotected copy.",
                            code="resume_encrypted",
                        ) from exc

                chunks = []
                for page in reader.pages:
                    try:
                        page_text = page.extract_text()
                    except Exception:  # pragma: no cover - per-page failure
                        page_text = None
                    if page_text:
                        chunks.append(page_text)
                text = "\n".join(chunks)

            elif ext == "docx":
                # python-docx reliably accepts file-like objects such as Django
                # UploadedFile. Using it here avoids docx2txt path/file-object
                # differences that can make uploads fail in production.
                from docx import Document

                document = Document(file_obj)
                chunks = [p.text for p in document.paragraphs if p.text]
                for table in document.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            chunks.append(" | ".join(cells))
                text = "\n".join(chunks)

            else:  # txt
                raw = file_obj.read()
                if isinstance(raw, bytes):
                    text = raw.decode("utf-8", errors="ignore")
                else:
                    text = str(raw)

        except ResumeParseError:
            raise
        except Exception as exc:
            logger.error("Error parsing resume '%s': %s", filename, exc)
            raise ResumeParseError(
                "The resume file could not be read. It may be corrupted or "
                "image-only. Try exporting it again as a text-based PDF or DOCX.",
                code="resume_unreadable",
            ) from exc

        text = (text or "").strip()

        if not text:
            raise ResumeParseError(
                "No readable text was found in the resume. Scanned or image-only "
                "documents are not supported - upload a text-based PDF or DOCX.",
                code="resume_no_text",
            )

        # Reset the file pointer so callers can safely reuse the file object
        # (e.g. Django saving it to a FileField after text extraction).
        try:
            file_obj.seek(0)
        except (OSError, AttributeError):
            pass

        return text
